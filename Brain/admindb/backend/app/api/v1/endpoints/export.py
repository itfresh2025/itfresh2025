import io
from typing import Optional

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ....database import get_db
from ....models.client import Client, ColorMark
from ....models.user import User
from ....api.deps import get_current_user
from ....core.encryption import decrypt_password

router = APIRouter()

_FILLS = {
    "none":   "FFFFFF",
    "green":  "CCFFCC",
    "yellow": "FFFACC",
    "red":    "FFCCCC",
}


@router.get("/clients/excel")
def export_excel(
    search: Optional[str] = Query(None),
    color_mark: Optional[ColorMark] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    q = db.query(Client).filter(Client.is_active == True)
    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            Client.company.ilike(like),
            Client.ip_address.ilike(like),
            Client.login.ilike(like),
        ))
    if color_mark:
        q = q.filter(Client.color_mark == color_mark)

    clients = q.order_by(Client.company).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Клиенты"

    headers = ["ID", "Фирма", "IP-адрес", "Маска", "Шлюз", "Логин",
               "Пароль", "TeamViewer ID", "Дата", "Метка", "Ссылка", "Заметки"]

    hdr_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

    for row_i, c in enumerate(clients, 2):
        pw = decrypt_password(c.password_encrypted) if c.password_encrypted else ""
        values = [
            c.id, c.company, c.ip_address, c.mask, c.gate, c.login, pw,
            c.teamviewer_id, str(c.date) if c.date else "",
            c.color_mark.value, c.info_link, c.notes,
        ]
        row_fill = PatternFill(
            start_color=_FILLS.get(c.color_mark.value, "FFFFFF"),
            end_color=_FILLS.get(c.color_mark.value, "FFFFFF"),
            fill_type="solid",
        )
        for col_i, val in enumerate(values, 1):
            cell = ws.cell(row=row_i, column=col_i, value=val or "")
            cell.fill = row_fill

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 45)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=admindb_clients.xlsx"},
    )
