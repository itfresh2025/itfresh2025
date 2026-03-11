from typing import List, Optional
from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel

from ....database import get_db
from ....models.tickets import Ticket, TicketStatus, TicketPriority, TicketSource
from ....models.ticket_comments import TicketComment
from ....models.ticket_events import TicketEvent, EventType
from ....models.ticket_relations import TicketRelation, RelationType
from ....models.client import Client
from ....models.user import User
from ....api.deps import get_current_user, require_admin

router = APIRouter()

# SLA limits (minutes)
SLA_MINUTES = {
    TicketPriority.critical: 60,
    TicketPriority.high: 240,
    TicketPriority.medium: 1440,
    TicketPriority.low: 4320,
}


def _sla_remaining(ticket: Ticket) -> Optional[int]:
    limit = SLA_MINUTES.get(ticket.priority)
    if limit is None:
        return None
    if ticket.status in (TicketStatus.resolved, TicketStatus.closed):
        return None
    if ticket.sla_paused:
        return None
    now = datetime.now(timezone.utc)
    created = ticket.created_at
    if created is None:
        return None
    if created.tzinfo is None:
        created = created.replace(tzinfo=timezone.utc)
    elapsed = int((now - created).total_seconds() // 60)
    return limit - elapsed  # negative means overdue


def _compute_overdue(ticket: Ticket) -> bool:
    remaining = _sla_remaining(ticket)
    if remaining is None:
        return False
    return remaining < 0


# ─── Schemas ──────────────────────────────────────────────────────────────────

class TicketCommentOut(BaseModel):
    id: int
    ticket_id: int
    author_id: int
    author_username: Optional[str] = None
    text: str
    created_at: str
    telegram_chat_id: Optional[str] = None

    class Config:
        from_attributes = True


class TicketOut(BaseModel):
    id: int
    client_id: Optional[int] = None
    client_name: Optional[str] = None
    title: str
    description: Optional[str] = None
    status: TicketStatus
    priority: TicketPriority
    created_by_id: int
    created_by_username: Optional[str] = None
    assigned_to_id: Optional[int] = None
    assigned_to_username: Optional[str] = None
    created_at: str
    updated_at: Optional[str] = None
    resolved_at: Optional[str] = None
    assigned_at: Optional[str] = None
    pending_reason: Optional[str] = None
    sla_paused: bool = False
    source: TicketSource
    telegram_chat_id: Optional[str] = None
    telegram_message_id: Optional[str] = None
    is_overdue: bool = False
    sla_remaining_minutes: Optional[int] = None
    comments_count: int = 0
    notify_channel_id: Optional[int] = None
    shared_fields: Optional[str] = '[]'

    class Config:
        from_attributes = True


class TicketCreate(BaseModel):
    client_id: Optional[int] = None
    title: str
    description: Optional[str] = None
    priority: TicketPriority = TicketPriority.medium
    assigned_to_id: Optional[int] = None
    source: TicketSource = TicketSource.manual
    telegram_chat_id: Optional[str] = None
    telegram_message_id: Optional[str] = None
    notify_channel_id: Optional[int] = None
    shared_fields: Optional[str] = '[]'


class TicketUpdate(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    status: Optional[TicketStatus] = None
    priority: Optional[TicketPriority] = None
    assigned_to_id: Optional[int] = None
    client_id: Optional[int] = None
    pending_reason: Optional[str] = None
    sla_paused: Optional[bool] = None
    notify_channel_id: Optional[int] = None
    shared_fields: Optional[str] = None


class TicketCommentCreate(BaseModel):
    text: str
    telegram_chat_id: Optional[str] = None


class TicketStatsOut(BaseModel):
    new: int = 0
    assigned: int = 0
    in_progress: int = 0
    pending: int = 0
    resolved: int
    closed: int
    total: int
    overdue: int


class TicketEventOut(BaseModel):
    id: int
    ticket_id: int
    user_id: Optional[int] = None
    user_username: Optional[str] = None
    event_type: EventType
    old_value: Optional[str] = None
    new_value: Optional[str] = None
    comment: Optional[str] = None
    created_at: str

    class Config:
        from_attributes = True


class TicketRelationCreate(BaseModel):
    related_ticket_id: int
    relation_type: RelationType = RelationType.related


class TicketRelationOut(BaseModel):
    id: int
    ticket_id: int
    related_ticket_id: int
    relation_type: RelationType
    related_ticket_title: Optional[str] = None

    class Config:
        from_attributes = True


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _ticket_to_out(t: Ticket, db: Session) -> TicketOut:
    return TicketOut(
        id=t.id,
        client_id=t.client_id,
        client_name=t.client.company if t.client else None,
        title=t.title,
        description=t.description,
        status=t.status,
        priority=t.priority,
        created_by_id=t.created_by_id,
        created_by_username=t.created_by.username if t.created_by else None,
        assigned_to_id=t.assigned_to_id,
        assigned_to_username=t.assigned_to.username if t.assigned_to else None,
        created_at=t.created_at.isoformat() if t.created_at else "",
        updated_at=t.updated_at.isoformat() if t.updated_at else None,
        resolved_at=t.resolved_at.isoformat() if t.resolved_at else None,
        assigned_at=t.assigned_at.isoformat() if t.assigned_at else None,
        pending_reason=t.pending_reason,
        sla_paused=t.sla_paused or False,
        source=t.source,
        telegram_chat_id=t.telegram_chat_id,
        telegram_message_id=t.telegram_message_id,
        is_overdue=_compute_overdue(t),
        sla_remaining_minutes=_sla_remaining(t),
        comments_count=len(t.comments),
        notify_channel_id=getattr(t, 'notify_channel_id', None),
        shared_fields=getattr(t, 'shared_fields', '[]') or '[]',
    )


def _write_event(
    db: Session,
    ticket_id: int,
    user_id: Optional[int],
    event_type: EventType,
    old_value: Optional[str] = None,
    new_value: Optional[str] = None,
    comment: Optional[str] = None,
):
    event = TicketEvent(
        ticket_id=ticket_id,
        user_id=user_id,
        event_type=event_type,
        old_value=old_value,
        new_value=new_value,
        comment=comment,
    )
    db.add(event)


# ─── Endpoints ────────────────────────────────────────────────────────────────

@router.get("/stats", response_model=TicketStatsOut)
def get_ticket_stats(
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    tickets = db.query(Ticket).all()
    return TicketStatsOut(
        new=sum(1 for t in tickets if t.status == TicketStatus.new),
        assigned=sum(1 for t in tickets if t.status == TicketStatus.assigned),
        in_progress=sum(1 for t in tickets if t.status == TicketStatus.in_progress),
        pending=sum(1 for t in tickets if t.status == TicketStatus.pending),
        resolved=sum(1 for t in tickets if t.status == TicketStatus.resolved),
        closed=sum(1 for t in tickets if t.status == TicketStatus.closed),
        total=len(tickets),
        overdue=sum(1 for t in tickets if _compute_overdue(t)),
    )


@router.get("/export")
def export_tickets(
    status: Optional[TicketStatus] = Query(None),
    priority: Optional[TicketPriority] = Query(None),
    client_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    """Экспорт тикетов в Excel (openpyxl)."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl not installed. Run: pip install openpyxl",
        )

    q = db.query(Ticket)
    if status:
        q = q.filter(Ticket.status == status)
    if priority:
        q = q.filter(Ticket.priority == priority)
    if client_id:
        q = q.filter(Ticket.client_id == client_id)

    tickets = q.order_by(Ticket.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tickets"

    headers = [
        "ID", "Клиент", "Заголовок", "Статус", "Приоритет",
        "Создан", "Назначен", "Создан_at", "Обновлён_at", "Закрыт_at",
        "SLA_осталось_мин", "Источник",
    ]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, t in enumerate(tickets, 2):
        remaining = _sla_remaining(t)
        ws.cell(row=row_idx, column=1, value=t.id)
        ws.cell(row=row_idx, column=2, value=t.client.company if t.client else "")
        ws.cell(row=row_idx, column=3, value=t.title)
        ws.cell(row=row_idx, column=4, value=t.status.value)
        ws.cell(row=row_idx, column=5, value=t.priority.value)
        ws.cell(row=row_idx, column=6, value=t.created_by.username if t.created_by else "")
        ws.cell(row=row_idx, column=7, value=t.assigned_to.username if t.assigned_to else "")
        ws.cell(row=row_idx, column=8, value=t.created_at.isoformat() if t.created_at else "")
        ws.cell(row=row_idx, column=9, value=t.updated_at.isoformat() if t.updated_at else "")
        ws.cell(row=row_idx, column=10, value=t.resolved_at.isoformat() if t.resolved_at else "")
        ws.cell(row=row_idx, column=11, value=remaining)
        ws.cell(row=row_idx, column=12, value=t.source.value)

        # Подсветка просроченных
        if remaining is not None and remaining < 0:
            red_fill = PatternFill(start_color="FFB3B3", end_color="FFB3B3", fill_type="solid")
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = red_fill

    # Авто-ширина колонок
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=tickets.xlsx"},
    )


@router.get("/", response_model=List[TicketOut])
def list_tickets(
    status: Optional[TicketStatus] = Query(None),
    priority: Optional[TicketPriority] = Query(None),
    client_id: Optional[int] = Query(None),
    assigned_to: Optional[int] = Query(None),
    my_only: bool = Query(False),
    overdue_only: bool = Query(False),
    search: Optional[str] = Query(None),
    skip: int = 0,
    limit: int = 200,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    q = db.query(Ticket)
    if status:
        q = q.filter(Ticket.status == status)
    if priority:
        q = q.filter(Ticket.priority == priority)
    if client_id:
        q = q.filter(Ticket.client_id == client_id)
    if assigned_to:
        q = q.filter(Ticket.assigned_to_id == assigned_to)
    if my_only:
        q = q.filter(Ticket.assigned_to_id == current_user.id)
    if search:
        pattern = f"%{search}%"
        q = q.filter(Ticket.title.ilike(pattern) | Ticket.description.ilike(pattern))

    tickets = q.order_by(Ticket.created_at.desc()).offset(skip).limit(limit).all()

    if overdue_only:
        tickets = [t for t in tickets if _compute_overdue(t)]

    return [_ticket_to_out(t, db) for t in tickets]


@router.post("/", response_model=TicketOut)
def create_ticket(
    data: TicketCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if data.client_id:
        client = db.query(Client).filter(Client.id == data.client_id).first()
        if not client:
            raise HTTPException(status_code=404, detail="Client not found")

    ticket = Ticket(
        client_id=data.client_id,
        title=data.title,
        description=data.description,
        priority=data.priority,
        assigned_to_id=data.assigned_to_id,
        created_by_id=current_user.id,
        source=data.source,
        telegram_chat_id=data.telegram_chat_id,
        telegram_message_id=data.telegram_message_id,
        notify_channel_id=data.notify_channel_id,
        shared_fields=data.shared_fields or '[]',
    )
    if data.assigned_to_id:
        ticket.status = TicketStatus.assigned
        ticket.assigned_at = datetime.now(timezone.utc)

    db.add(ticket)
    db.flush()  # get ticket.id before writing event

    _write_event(
        db,
        ticket_id=ticket.id,
        user_id=current_user.id,
        event_type=EventType.created,
        new_value=ticket.status.value,
    )

    db.commit()
    db.refresh(ticket)

    # Send Telegram notification
    try:
        from ....bot.telegram_bot import notify_new_ticket
        import asyncio
        asyncio.create_task(notify_new_ticket(ticket.id))
    except Exception:
        pass

    return _ticket_to_out(ticket, db)


@router.get("/{ticket_id}", response_model=TicketOut)
def get_ticket(
    ticket_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    return _ticket_to_out(ticket, db)


@router.patch("/{ticket_id}", response_model=TicketOut)
def update_ticket(
    ticket_id: int,
    data: TicketUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    update_data = data.model_dump(exclude_unset=True)

    # Track status change
    old_status = ticket.status
    old_priority = ticket.priority
    old_assigned = ticket.assigned_to_id

    for field, value in update_data.items():
        setattr(ticket, field, value)

    now = datetime.now(timezone.utc)

    # Set resolved_at when status changes to resolved/closed
    if "status" in update_data and update_data["status"] in (
        TicketStatus.resolved.value, TicketStatus.closed.value,
        "resolved", "closed",
    ):
        if not ticket.resolved_at:
            ticket.resolved_at = now

    # Set assigned_at when assigned_to changes
    if "assigned_to_id" in update_data and update_data["assigned_to_id"] and \
            update_data["assigned_to_id"] != old_assigned:
        ticket.assigned_at = now
        if ticket.status == TicketStatus.new:
            ticket.status = TicketStatus.assigned

    # SLA pause on pending
    if "status" in update_data:
        new_status = ticket.status
        if new_status == TicketStatus.pending:
            ticket.sla_paused = True
        elif old_status == TicketStatus.pending and new_status != TicketStatus.pending:
            ticket.sla_paused = False

    # Write events
    if "status" in update_data and ticket.status != old_status:
        _write_event(
            db,
            ticket_id=ticket.id,
            user_id=current_user.id,
            event_type=EventType.status_changed,
            old_value=old_status.value,
            new_value=ticket.status.value,
        )

    if "priority" in update_data and ticket.priority != old_priority:
        _write_event(
            db,
            ticket_id=ticket.id,
            user_id=current_user.id,
            event_type=EventType.priority_changed,
            old_value=old_priority.value,
            new_value=ticket.priority.value,
        )

    if "assigned_to_id" in update_data and ticket.assigned_to_id != old_assigned:
        _write_event(
            db,
            ticket_id=ticket.id,
            user_id=current_user.id,
            event_type=EventType.assigned,
            old_value=str(old_assigned) if old_assigned else None,
            new_value=str(ticket.assigned_to_id) if ticket.assigned_to_id else None,
        )

    db.commit()
    db.refresh(ticket)
    return _ticket_to_out(ticket, db)


@router.delete("/{ticket_id}")
def delete_ticket(
    ticket_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    db.delete(ticket)
    db.commit()
    return {"ok": True}


@router.get("/{ticket_id}/comments", response_model=List[TicketCommentOut])
def get_comments(
    ticket_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    return [
        TicketCommentOut(
            id=c.id,
            ticket_id=c.ticket_id,
            author_id=c.author_id,
            author_username=c.author.username if c.author else None,
            text=c.text,
            created_at=c.created_at.isoformat() if c.created_at else "",
            telegram_chat_id=c.telegram_chat_id,
        )
        for c in ticket.comments
    ]


@router.post("/{ticket_id}/comments", response_model=TicketCommentOut)
def add_comment(
    ticket_id: int,
    data: TicketCommentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    comment = TicketComment(
        ticket_id=ticket_id,
        author_id=current_user.id,
        text=data.text,
        telegram_chat_id=data.telegram_chat_id,
    )
    db.add(comment)
    db.flush()

    _write_event(
        db,
        ticket_id=ticket_id,
        user_id=current_user.id,
        event_type=EventType.comment_added,
        comment=data.text[:200],
    )

    db.commit()
    db.refresh(comment)

    return TicketCommentOut(
        id=comment.id,
        ticket_id=comment.ticket_id,
        author_id=comment.author_id,
        author_username=current_user.username,
        text=comment.text,
        created_at=comment.created_at.isoformat() if comment.created_at else "",
        telegram_chat_id=comment.telegram_chat_id,
    )


@router.get("/{ticket_id}/events", response_model=List[TicketEventOut])
def get_ticket_events(
    ticket_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    return [
        TicketEventOut(
            id=e.id,
            ticket_id=e.ticket_id,
            user_id=e.user_id,
            user_username=e.user.username if e.user else None,
            event_type=e.event_type,
            old_value=e.old_value,
            new_value=e.new_value,
            comment=e.comment,
            created_at=e.created_at.isoformat() if e.created_at else "",
        )
        for e in sorted(ticket.events, key=lambda x: x.created_at or datetime.min)
    ]


@router.post("/{ticket_id}/relations", response_model=TicketRelationOut)
def create_relation(
    ticket_id: int,
    data: TicketRelationCreate,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    related = db.query(Ticket).filter(Ticket.id == data.related_ticket_id).first()
    if not related:
        raise HTTPException(status_code=404, detail="Related ticket not found")

    if ticket_id == data.related_ticket_id:
        raise HTTPException(status_code=400, detail="Cannot relate ticket to itself")

    existing = db.query(TicketRelation).filter(
        TicketRelation.ticket_id == ticket_id,
        TicketRelation.related_ticket_id == data.related_ticket_id,
    ).first()
    if existing:
        raise HTTPException(status_code=409, detail="Relation already exists")

    relation = TicketRelation(
        ticket_id=ticket_id,
        related_ticket_id=data.related_ticket_id,
        relation_type=data.relation_type,
    )
    db.add(relation)
    db.commit()
    db.refresh(relation)

    return TicketRelationOut(
        id=relation.id,
        ticket_id=relation.ticket_id,
        related_ticket_id=relation.related_ticket_id,
        relation_type=relation.relation_type,
        related_ticket_title=related.title,
    )


@router.get("/{ticket_id}/relations", response_model=List[TicketRelationOut])
def get_relations(
    ticket_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    return [
        TicketRelationOut(
            id=r.id,
            ticket_id=r.ticket_id,
            related_ticket_id=r.related_ticket_id,
            relation_type=r.relation_type,
            related_ticket_title=r.related_ticket.title if r.related_ticket else None,
        )
        for r in ticket.relations
    ]
